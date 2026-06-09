import os
import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from openpyxl.styles import Font
import traceback

class VirtualAggregationApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Виртуальная агрегация (Excel)")
        self.root.geometry("500x550")
        self.root.resizable(False, False)

        self.parent_file = ""
        self.child_files = []

        # UI Элементы
        frame = tk.Frame(root, padx=20, pady=20)
        frame.pack(expand=True, fill="both")

        tk.Label(frame, text="1. Файл с кодами наборов (родители):", font=("Arial", 10, "bold")).pack(anchor="w")
        self.btn_parent = tk.Button(frame, text="Выбрать файл", command=self.select_parent)
        self.btn_parent.pack(fill="x", pady=(5, 5))
        self.lbl_parent = tk.Label(frame, text="Файл не выбран", fg="gray", wraplength=450)
        self.lbl_parent.pack(anchor="w", pady=(0, 15))

        tk.Label(frame, text="2. Кол-во вложений из КАЖДОГО файла на 1 набор:", font=("Arial", 10, "bold")).pack(anchor="w")
        self.ent_count = tk.Entry(frame)
        self.ent_count.insert(0, "1")
        self.ent_count.pack(fill="x", pady=(5, 15))

        tk.Label(frame, text="3. Файлы с кодами вложений (дети):", font=("Arial", 10, "bold")).pack(anchor="w")
        self.btn_children = tk.Button(frame, text="Выбрать файлы", command=self.select_children)
        self.btn_children.pack(fill="x", pady=(5, 5))
        self.lbl_children = tk.Label(frame, text="Файлы не выбраны", fg="gray", wraplength=450)
        self.lbl_children.pack(anchor="w", pady=(0, 20))

        self.btn_process = tk.Button(
            frame,
            text="Выполнить агрегацию в Excel",
            bg="#2196F3",
            fg="white",
            font=("Arial", 12, "bold"),
            command=self.process
        )
        self.btn_process.pack(fill="x", pady=(10, 0), ipady=10)

        tk.Label(frame, text="Результат: один файл Excel (.xlsx)", font=("Arial", 8), fg="blue").pack(pady=(10, 0))

    def select_parent(self):
        file = filedialog.askopenfilename(
            title="Выберите файл с кодами наборов",
            filetypes=[("TXT files", "*.txt"), ("All files", "*.*")]
        )
        if file:
            self.parent_file = file
            self.lbl_parent.config(text=os.path.basename(file), fg="black")

    def select_children(self):
        files = filedialog.askopenfilenames(
            title="Выберите файлы с кодами вложений",
            filetypes=[("TXT files", "*.txt"), ("All files", "*.*")]
        )
        if files:
            self.child_files = list(files)
            self.lbl_children.config(text=f"Выбрано файлов: {len(files)}", fg="black")

    def read_codes(self, filepath):
        codes = []
        for encoding in ['utf-8-sig', 'utf-8', 'cp1251', 'utf-16']:
            try:
                with open(filepath, 'r', encoding=encoding) as f:
                    for line in f:
                        code = line.strip()
                        if code:
                            codes.append(code)
                if codes: return codes
            except:
                continue
        return codes

    def process(self):
        try:
            if not self.parent_file:
                messagebox.showerror("Ошибка", "Выберите файл с кодами наборов")
                return

            try:
                count_per_file = int(self.ent_count.get())
                if count_per_file <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Ошибка", "Введите корректное число вложений (целое, больше 0)")
                return

            if not self.child_files:
                messagebox.showerror("Ошибка", "Выберите файлы с кодами вложений")
                return

            # Чтение данных
            parent_codes = self.read_codes(self.parent_file)
            if not parent_codes:
                messagebox.showerror("Ошибка", f"Файл родителей пуст или не прочитан: {os.path.basename(self.parent_file)}")
                return

            child_files_data = []
            for cf in self.child_files:
                codes = self.read_codes(cf)
                if codes:
                    child_files_data.append(codes)

            if not child_files_data:
                messagebox.showerror("Ошибка", "Нет данных для вложений")
                return

            # Расчет количества наборов
            min_attachments = min(len(c_list) // count_per_file for c_list in child_files_data)
            actual_sets = min(len(parent_codes), min_attachments)

            if actual_sets == 0:
                messagebox.showerror("Ошибка", "Недостаточно данных для создания наборов.\nПроверьте количество кодов в файлах.")
                return

            # Выбор пути сохранения
            save_path = filedialog.asksaveasfilename(
                title="Сохранить как...",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile="Результат_агрегации"
            )

            if not save_path:
                return

            # Создание Excel
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Агрегация"

            # Заголовки (используем русские названия согласно запросу)
            ws.append(["Код набора", "Код вложения"])
            ws['A1'].font = Font(bold=True)
            ws['B1'].font = Font(bold=True)

            row_idx = 2
            for i in range(actual_sets):
                p_code = parent_codes[i]
                for codes_list in child_files_data:
                    start_idx = i * count_per_file
                    for k in range(count_per_file):
                        c_code = codes_list[start_idx + k]
                        ws.cell(row=row_idx, column=1, value=p_code)
                        ws.cell(row=row_idx, column=2, value=c_code)
                        row_idx += 1

            wb.save(save_path)
            messagebox.showinfo("Успех", f"Агрегация завершена!\n\nФайл сохранен: {save_path}\nСоздано наборов: {actual_sets}")

        except PermissionError:
            messagebox.showerror("Ошибка", "Не удалось сохранить файл. Возможно, он открыт в Excel.")
        except Exception as e:
            error_msg = traceback.format_exc()
            messagebox.showerror("Ошибка", f"Произошла ошибка:\n{e}\n\n{error_msg}")

if __name__ == "__main__":
    root = tk.Tk()
    app = VirtualAggregationApp(root)
    root.mainloop()
